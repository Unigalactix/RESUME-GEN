import io
import unittest

from docx import Document

from tools.resume_scorer import extract_resume_text


class FakeUpload(io.BytesIO):
    def __init__(self, data, name):
        super().__init__(data)
        self.name = name

    def getvalue(self):
        return super().getvalue()


class ResumeUploadFormatTests(unittest.TestCase):
    def test_extract_resume_text_supports_txt(self):
        uploaded = FakeUpload(b"Alex Chen\nPython\nAWS\n", "resume.txt")

        text = extract_resume_text(uploaded)

        self.assertIn("Alex Chen", text)
        self.assertIn("Python", text)

    def test_extract_resume_text_supports_docx(self):
        document = Document()
        document.add_paragraph("Alex Chen")
        document.add_paragraph("Built backend systems in Python and AWS.")
        buffer = io.BytesIO()
        document.save(buffer)
        uploaded = FakeUpload(buffer.getvalue(), "resume.docx")

        text = extract_resume_text(uploaded)

        self.assertIn("Alex Chen", text)
        self.assertIn("backend systems", text)

    def test_extract_resume_text_supports_doc_best_effort(self):
        fake_doc_bytes = b"\xd0\xcf\x11\xe0" + "Alex Chen Resume".encode("utf-16le") + b"Python AWS Backend Engineer"
        uploaded = FakeUpload(fake_doc_bytes, "resume.doc")

        text = extract_resume_text(uploaded)

        self.assertIn("Alex Chen Resume", text)
        self.assertIn("Python AWS Backend Engineer", text)


if __name__ == "__main__":
    unittest.main()